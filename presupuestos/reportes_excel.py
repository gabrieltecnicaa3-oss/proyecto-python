"""Reportes Excel del módulo Presupuestos (sección 5.2 y 5.3 de la
especificación), generados con openpyxl.

No recalcula nada de cero: arma los workbooks a partir de los totales que
ya devuelve el motor de cálculo (calculo_presupuesto.py), vía
`calcular_totales_por_concepto` / `calcular_explosion_insumos` /
`calcular_reporte_prevision_fondos`. Acá solo vive el armado del Excel
(hojas, columnas, formato) y la resolución de nombres (tareas, obra).
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .calculo_presupuesto import (
    calcular_totales_por_concepto,
    calcular_explosion_insumos,
    calcular_reporte_prevision_fondos,
)
from .models import listar_tareas, listar_categorias_odoo
from .routes import _calcular_resultado_tarea

_FONT_HEADER = Font(bold=True, color="FFFFFF")
_FILL_HEADER = PatternFill("solid", fgColor="1F2937")
_FONT_TOTAL = Font(bold=True)
_FILL_TOTAL = PatternFill("solid", fgColor="EEF2FF")

_GUION = "—"


def _escribir_tabla_explosion(ws, filas):
    """Escribe la tabla `Categoría | Fabricación ($) | Montaje ($)` de la
    sección 5.2 (usada tanto en la pestaña de cada tarea como en "Resumen")."""
    ws.append(["Categoría", "Fabricación ($)", "Montaje ($)"])
    for celda in ws[1]:
        celda.font = _FONT_HEADER
        celda.fill = _FILL_HEADER

    for fila in filas:
        es_total = fila["categoria"] == "TOTAL"
        valores = [
            fila["categoria"],
            _GUION if fila["fabricacion"] is None else round(fila["fabricacion"], 2),
            _GUION if fila["montaje"] is None else round(fila["montaje"], 2),
        ]
        ws.append(valores)
        if es_total:
            for celda in ws[ws.max_row]:
                celda.font = _FONT_TOTAL
                celda.fill = _FILL_TOTAL

    for col in range(1, 4):
        ws.column_dimensions[get_column_letter(col)].width = 28 if col == 1 else 20


def _nombres_equipos(db, equipo_ids):
    """Resuelve {equipo_id: nombre} desde catalogo_equipos (sin FK dura:
    si la tabla no existe todavía, devuelve vacío), mismo criterio que
    `_resumen_recursos_con_nombres` en routes.py."""
    ids = sorted({int(eid) for eid in equipo_ids if eid})
    if not ids:
        return {}
    placeholders = ",".join("?" for _ in ids)
    try:
        rows = db.execute(
            f"SELECT id, nombre FROM catalogo_equipos WHERE id IN ({placeholders})", tuple(ids)
        ).fetchall()
    except Exception:
        return {}
    return {r[0]: r[1] for r in rows}


def _escribir_tabla_recursos(ws, db, resultados):
    """Agrupa por recurso; el unitario es total dividido por cantidad."""
    recursos = {}
    for tarea in resultados:
        for seccion in ("fabricacion", "montaje"):
            for item in tarea[seccion]["items"]:
                rubro = item["rubro"]
                datos = item.get("datos") or {}
                equipo_id = None
                if rubro == "mano_obra":
                    nombre = "Fabricación" if seccion == "fabricacion" else "Montaje"
                    unidad = "Operarios-día"
                    cantidad = float(datos.get("operarios") or 0) * float(datos.get("dias") or 0)
                elif seccion == "montaje" and rubro in ("equipo", "ingeniero", "tecnico_hys"):
                    nombre = {"equipo": "Equipo", "ingeniero": "Director de obra",
                              "tecnico_hys": "Técnico H y S"}[rubro]
                    unidad = "Días"
                    cantidad = float(datos.get("dias") or 0)
                    if rubro == "equipo":
                        equipo_id = datos.get("equipo_id")
                else:
                    continue
                if cantidad == 0 and float(item["subtotal"]) == 0:
                    continue
                clave = (seccion, rubro, equipo_id)
                entrada = recursos.setdefault(clave, {
                    "nombre": nombre, "unidad": unidad, "equipo_id": equipo_id,
                    "rubro": rubro, "cantidad": 0.0, "total": 0.0,
                })
                entrada["cantidad"] += cantidad
                entrada["total"] += float(item["subtotal"])

    nombres_equipo = _nombres_equipos(
        db, (r["equipo_id"] for r in recursos.values() if r["rubro"] == "equipo")
    )
    ws.append([])
    ws.append(["RECURSOS"])
    ws[ws.max_row][0].font = _FONT_HEADER
    ws[ws.max_row][0].fill = _FILL_HEADER

    ws.append(["Recurso", "Cantidad", "Unidad", "Precio unitario ($)", "Total ($)"])
    for celda in ws[ws.max_row]:
        celda.font = _FONT_TOTAL
        celda.fill = _FILL_TOTAL
    if recursos:
        for recurso in recursos.values():
            nombre = recurso["nombre"]
            if recurso["rubro"] == "equipo":
                equipo_id = recurso["equipo_id"]
                nombre = nombres_equipo.get(equipo_id) or (
                    f"Equipo #{equipo_id}" if equipo_id else "(sin equipo asociado)"
                )
            ws.append([nombre, recurso["cantidad"], recurso["unidad"],
                       recurso["total"] / recurso["cantidad"] if recurso["cantidad"] else None,
                       recurso["total"]])
            ws.cell(ws.max_row, 2).number_format = "0.00"
            for col in (4, 5):
                ws.cell(ws.max_row, col).number_format = '"$" #,##0.00'
        ws.append(["TOTAL RECURSOS", None, None, None,
                   sum(r["total"] for r in recursos.values())])
        for celda in ws[ws.max_row]:
            celda.font = _FONT_TOTAL
            celda.fill = _FILL_TOTAL
        ws.cell(ws.max_row, 5).number_format = '"$" #,##0.00'
    else:
        ws.append(["Sin recursos cargados."])
    for col in (3, 4, 5):
        ws.column_dimensions[get_column_letter(col)].width = 22


def generar_reporte_explosion_insumos(db, presupuesto_id):
    """Sección 5.2 — Reporte 1: una pestaña por tarea + pestaña "Resumen"."""
    tareas = listar_tareas(db, presupuesto_id)

    wb = Workbook()
    wb.remove(wb.active)

    resultados_todas_tareas = []
    for tarea in tareas:
        resultado = _calcular_resultado_tarea(db, tarea["id"])
        resultados_todas_tareas.append(resultado)

        totales = calcular_totales_por_concepto([resultado])
        explosion = calcular_explosion_insumos(totales)

        nombre_hoja = str(tarea["nombre"])[:31] or f"Tarea {tarea['id']}"
        # Nombres de hoja repetidos/ inválidos (Excel no permite : \ / ? * [ ]).
        for caracter in '[]:*?/\\':
            nombre_hoja = nombre_hoja.replace(caracter, "-")
        nombre_base, sufijo = nombre_hoja, 2
        nombres_usados = {ws.title for ws in wb.worksheets}
        while nombre_hoja in nombres_usados:
            nombre_hoja = f"{nombre_base[:28]} ({sufijo})"
            sufijo += 1

        ws = wb.create_sheet(title=nombre_hoja)
        _escribir_tabla_explosion(ws, explosion["filas"])
        _escribir_tabla_recursos(ws, db, [resultado])

    totales_presupuesto = calcular_totales_por_concepto(resultados_todas_tareas)
    explosion_total = calcular_explosion_insumos(totales_presupuesto)
    ws_resumen = wb.create_sheet(title="Resumen", index=0)
    _escribir_tabla_explosion(ws_resumen, explosion_total["filas"])
    _escribir_tabla_recursos(ws_resumen, db, resultados_todas_tareas)
    wb.active = 0

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_reporte_prevision_fondos(db, presupuesto_id, obra_referencia):
    """Sección 5.3 — Reporte 2: una sola tabla reclasificada por
    categoria_odoo (config_categorias_odoo, sección 2.1)."""
    tareas = listar_tareas(db, presupuesto_id)
    resultados = [_calcular_resultado_tarea(db, tarea["id"]) for tarea in tareas]
    totales = calcular_totales_por_concepto(resultados)
    mapeo = listar_categorias_odoo(db)
    reporte = calcular_reporte_prevision_fondos(totales, mapeo)

    wb = Workbook()
    ws = wb.active
    ws.title = "Previsión de fondos"

    obra = obra_referencia or ""
    ws.append(["CATEGORIA", f"TA-{obra}", obra])
    for celda in ws[1]:
        celda.font = _FONT_HEADER
        celda.fill = _FILL_HEADER

    for fila in reporte["filas"]:
        fab = fila["fabricacion"]
        mon = fila["montaje"]
        ws.append([
            fila["categoria"],
            _GUION if fab is None else round(fab, 2),
            _GUION if mon is None else round(mon, 2),
        ])

    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
