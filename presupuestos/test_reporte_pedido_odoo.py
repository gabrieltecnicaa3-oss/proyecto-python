"""Tests del Reporte 3 (pedido "día 0" de Odoo). Mismo estilo que
test_calculo_presupuesto.py: asserts simples, se corre con
`python -m presupuestos.test_reporte_pedido_odoo`."""
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from .constants import ODOO_PEDIDO_COLUMNAS
from .models import _SEED_PRODUCTOS_ODOO
from .reportes_odoo_pedido import (
    construir_pedido,
    generar_excel_pedido_odoo,
    validar_analitica_id,
    _order_deadline,
    _order_reference,
    nombre_archivo_pedido,
)


def _reglas(seccion):
    return [
        {"seccion": s, "concepto": c, "tipo_item": t, "familia": f, "tarea": ta, "producto_odoo": p}
        for (s, c, t, f, ta, p) in _SEED_PRODUCTOS_ODOO
        if s == seccion
    ]


def _item(rubro, subtotal, tipo_item=None, tarea=None, **datos):
    return {"rubro": rubro, "tipo_item": tipo_item, "datos": datos, "subtotal": subtotal, "tarea": tarea}


def _por_producto(pedido):
    return {l["producto"]: l for l in pedido["lineas"]}


def test_fabricacion_agrupa_por_producto_y_excluye_rubros_no_cargables():
    items = [
        _item("materiales", 1000.0, "perfil", perfil_id=1),
        _item("materiales", 500.5, "perfil", perfil_id=2),
        _item("materiales", 100.0, "porcentaje"),
        _item("materiales", 300.0, "chapa", perfil_id=3),
        _item("bulones", 50.0),
        _item("pintura", 200.0),
        _item("ingenieria", 80.0),
        _item("subcontratos", 40.0),
        _item("mano_obra", 9999.0),
        _item("consumibles", 9999.0),
        _item("fletes", 9999.0),
    ]
    pedido = construir_pedido(items, _reglas("FABRICACION"), {1: "IPN", 2: "UPN", 3: "CH CINCALUM"}, {}, "FABRICACION")
    lineas = _por_producto(pedido)
    assert lineas["Perfiles Metálicos"]["precio_unitario"] == 1500.5
    assert lineas["Perfiles Metálicos"]["cantidad"] == 1
    assert lineas["Placas varias"]["precio_unitario"] == 100.0
    assert lineas["Chapas varias"]["precio_unitario"] == 300.0
    assert lineas["Bulones"]["precio_unitario"] == 50.0
    assert lineas["Pintura y Accesorios"]["precio_unitario"] == 200.0
    assert lineas["Proyectos de Ingeniería (EEMM)"]["precio_unitario"] == 80.0
    assert lineas["Subcontratos de Estructuras Metálicas y Herrería"]["precio_unitario"] == 40.0
    assert pedido["total_archivo"] == 2270.5
    assert pedido["total_rubros"] == 2270.5
    assert pedido["sin_mapear"] == []
    assert pedido["cuadra"]


def test_concepto_sin_producto_se_avisa_y_la_diferencia_cuadra():
    items = [
        _item("materiales", 1000.0, "perfil", perfil_id=1),
        _item("materiales", 75.25, "no_listado"),
    ]
    pedido = construir_pedido(items, _reglas("FABRICACION"), {1: "IPN"}, {}, "FABRICACION")
    assert pedido["sin_mapear"] == [{"concepto": "materiales / no_listado", "importe": 75.25}]
    assert pedido["total_archivo"] == 1000.0
    assert pedido["total_rubros"] == 1075.25
    assert pedido["diferencia"] == 75.25
    assert pedido["cuadra"]


def test_regla_por_familia_gana_y_producto_nulo_no_se_omite_en_silencio():
    reglas = _reglas("FABRICACION") + [
        {"seccion": "FABRICACION", "concepto": "materiales", "tipo_item": "perfil", "familia": "Correas", "producto_odoo": "Correas varias"},
        {"seccion": "FABRICACION", "concepto": "materiales", "tipo_item": "perfil", "familia": "Sin producto", "producto_odoo": None},
    ]
    items = [
        _item("materiales", 100.0, "perfil", perfil_id=1),
        _item("materiales", 200.0, "perfil", perfil_id=2),
        _item("materiales", 300.0, "perfil", perfil_id=3),
    ]
    pedido = construir_pedido(items, reglas, {1: "correas", 2: "IPN", 3: "Sin producto"}, {}, "FABRICACION")
    lineas = _por_producto(pedido)
    assert lineas["Correas varias"]["precio_unitario"] == 100.0
    assert lineas["Perfiles Metálicos"]["precio_unitario"] == 200.0
    assert pedido["sin_mapear"] == [{"concepto": "materiales / perfil (familia Sin producto)", "importe": 300.0}]


def test_montaje_equipos_por_producto_y_tarifa_mas_subcontratos():
    equipos = {
        1: {"id": 1, "nombre": "Hidrogrua", "producto_odoo": "Alquiler de HIDRO 42 TM"},
        2: {"id": 2, "nombre": "Brazo Articulado", "producto_odoo": None},
    }
    items = [
        _item("equipo", 5 * 1000.0, equipo_id=1, dias=5, tarifa_dia=1000.0),
        _item("equipo", 2 * 1000.0, equipo_id=1, dias=2, tarifa_dia=1000.0),
        _item("equipo", 3 * 1200.0, equipo_id=1, dias=3, tarifa_dia=1200.0),
        _item("equipo", 4 * 380.0, equipo_id=2, dias=4, tarifa_dia=380.0),
        _item("subcontratos", 90.0),
        _item("mano_obra", 5000.0),
        _item("ingeniero", 5000.0),
        _item("tecnico_hys", 5000.0),
    ]
    pedido = construir_pedido(items, _reglas("MONTAJE"), {}, equipos, "MONTAJE")
    hidro = [l for l in pedido["lineas"] if l["producto"] == "Alquiler de HIDRO 42 TM"]
    assert {(l["cantidad"], l["precio_unitario"]) for l in hidro} == {(7.0, 1000.0), (3.0, 1200.0)}
    sub = [l for l in pedido["lineas"] if l["producto"].startswith("Subcontratos")]
    assert sub[0]["cantidad"] == 1 and sub[0]["precio_unitario"] == 90.0
    assert pedido["sin_mapear"] == [{"concepto": "equipo / Brazo Articulado", "importe": 1520.0}]
    assert pedido["total_archivo"] == 7000.0 + 3600.0 + 90.0
    assert pedido["total_rubros"] == 7000.0 + 3600.0 + 1520.0 + 90.0
    assert pedido["diferencia"] == 1520.0 and pedido["cuadra"]


def test_importe_cero_sin_producto_no_genera_aviso():
    items = [_item("equipo", 0.0, equipo_id=9, dias=0, tarifa_dia=0)]
    pedido = construir_pedido(items, _reglas("MONTAJE"), {}, {}, "MONTAJE")
    assert pedido["sin_mapear"] == [] and pedido["lineas"] == []


def test_materiales_por_pestana_van_a_un_unico_producto():
    items = [
        _item("materiales", 100.0, "perfil", tarea="Correas", perfil_id=1),
        _item("materiales", 20.0, "porcentaje", tarea="Correas"),
        _item("bulones", 7.0, tarea="Correas"),
        _item("materiales", 300.0, "grating", tarea="Grating", perfil_id=2),
        _item("materiales", 30.0, "fijaciones", tarea="Grating"),
        _item("materiales", 400.0, "perfil", tarea="zinguería", perfil_id=1),
        _item("materiales", 50.0, "no_listado", tarea="Zinguería"),
        _item("materiales", 60.0, "perfil", tarea="Insertos", perfil_id=3),
        _item("materiales", 500.0, "perfil", tarea="Estructura Metálica", perfil_id=1),
        _item("materiales", 90.0, "chapa", tarea="Chapeado", perfil_id=4),
        _item("materiales", 11.0, "tornillos", tarea="Chapeado"),
    ]
    pedido = construir_pedido(items, _reglas("FABRICACION"), {1: "IPN", 2: "GRA Estandar", 3: "Redondo", 4: "CH CINCALUM"}, {}, "FABRICACION")
    lineas = {l["producto"]: l["precio_unitario"] for l in pedido["lineas"]}
    assert lineas == {
        "Correas varias": 120.0,
        "Bulones": 7.0,
        "Grating": 330.0,
        "Zinguerias varias": 450.0,
        "Varillas Roscadas": 60.0,
        "Perfiles Metálicos": 500.0,
        "Chapas varias": 90.0,
        "Consumibles varios": 11.0,
    }
    assert pedido["sin_mapear"] == [] and pedido["cuadra"]


def test_validar_analitica_id():
    assert validar_analitica_id(" 932 ") == 932
    for malo in ("", "   ", None, "0", "-5", "9.5", "abc", "12a", "1" * 13):
        try:
            validar_analitica_id(malo)
        except ValueError as e:
            assert str(e)
        else:
            raise AssertionError(f"debió rechazar {malo!r}")


def test_order_deadline_y_reference():
    assert _order_deadline("2026-03-05 18:30:00") == "2026-03-05 12:00:00"
    assert _order_deadline("2026-03-05") == "2026-03-05 12:00:00"
    assert _order_deadline(date(2026, 3, 5)) == "2026-03-05 12:00:00"
    assert _order_deadline(None, hoy=date(2026, 10, 9)) == "2026-10-09 12:00:00"
    assert _order_deadline("basura", hoy=date(2026, 10, 9)) == "2026-10-09 12:00:00"
    assert _order_reference("FABRICACION", "BUN-012") == "TA-BUN-012"
    assert _order_reference("MONTAJE", "BUN-012") == "BUN-012"
    try:
        _order_reference("MONTAJE", "  ")
    except ValueError:
        pass
    else:
        raise AssertionError("debió rechazar obra vacía")
    assert nombre_archivo_pedido("TA-BUN/012") == "odoo_pedido_TA-BUN_012.xlsx"


def test_excel_estructura_encabezados_cabecera_y_analitica():
    pedido = {
        "order_reference": "TA-BUN-012",
        "order_deadline": "2026-03-05 12:00:00",
        "lineas": [
            {"producto": "Perfiles Metálicos", "cantidad": 1, "precio_unitario": 1500.5},
            {"producto": "Bulones", "cantidad": 1, "precio_unitario": 50.0},
            {"producto": "Alquiler GRUA 30 TONELADAS", "cantidad": 7.5, "precio_unitario": 1000.0},
        ],
    }
    buffer = generar_excel_pedido_odoo(pedido, 932, vendor="Proveedor X")
    ws = load_workbook(BytesIO(buffer.getvalue())).active
    filas = [[c.value for c in fila] for fila in ws.iter_rows()]

    assert filas[0] == list(ODOO_PEDIDO_COLUMNAS)
    assert len(filas) == 4
    assert filas[1][:5] == ["TA-BUN-012", "Proveedor X", "2026-03-05 12:00:00", None, "Proveedor X"]
    assert filas[1][5:8] == ["Perfiles Metálicos", 1, 1500.5]
    for fila in filas[1:]:
        assert fila[:5] == ["TA-BUN-012", "Proveedor X", "2026-03-05 12:00:00", None, "Proveedor X"]
        assert fila[8] is None and fila[9] is None
        assert fila[10] == '{"932": 100}'
    assert filas[3][5:8] == ["Alquiler GRUA 30 TONELADAS", 7.5, 1000.0]


if __name__ == "__main__":
    pruebas = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    for nombre, funcion in pruebas:
        funcion()
        print(f"[OK] {nombre}")
    print(f"\n{len(pruebas)}/{len(pruebas)} tests OK")
