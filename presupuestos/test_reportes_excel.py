"""Regresiones de los recursos valorizados del reporte Excel."""
import sqlite3
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from .calculo_presupuesto import calcular_tarea
from .models import (
    ensure_tablas_presupuestos, crear_presupuesto, crear_tarea,
    crear_secciones_tarea, crear_item_costo, obtener_config,
)
from .reportes_excel import generar_reporte_explosion_insumos


class ReportesExcelTests(unittest.TestCase):
    def setUp(self):
        self.db = sqlite3.connect(":memory:")
        self.addCleanup(self.db.close)
        self.db.execute("CREATE TABLE catalogo_equipos (id INTEGER, nombre TEXT)")
        self.db.execute("INSERT INTO catalogo_equipos VALUES (7, 'Grua')")

    def _resultado(self, tarifa=1000):
        pcts = {"gg_pct": 0, "beneficio_pct": 0, "imp_pct": 0}
        return calcular_tarea([
            {"rubro": "mano_obra", "datos": {"operarios": 2, "dias": 3, "tarifa_dh": tarifa}},
            {"rubro": "consumibles", "datos": {"operarios": 2, "dias": 3, "tarifa_dh": 50}},
        ], pcts, [
            {"rubro": "mano_obra", "datos": {"operarios": 3, "dias": 2, "tarifa_dh": 2000}},
            {"rubro": "equipo", "datos": {"equipo_id": 7, "dias": 4, "tarifa_dia": 5000}},
            {"rubro": "ingeniero", "datos": {"dias": 5, "tarifa_dia": 3000}},
            {"rubro": "tecnico_hys", "datos": {"dias": 2, "tarifa_dia": 2500}},
        ], pcts, tipo_cambio_referencia=1500)

    def _workbook(self, resultados):
        tareas = [{"id": i + 1, "nombre": f"Tarea {i + 1}"} for i in range(len(resultados))]
        with patch("presupuestos.reportes_excel.listar_tareas", return_value=tareas), \
                patch("presupuestos.reportes_excel._calcular_resultado_tarea",
                      side_effect=resultados):
            buffer = generar_reporte_explosion_insumos(self.db, 1)
        wb = load_workbook(buffer)
        self.addCleanup(wb.close)
        return wb

    def _recursos(self, ws):
        filas = list(ws.iter_rows(values_only=True))
        inicio = next(i for i, fila in enumerate(filas) if fila[0] == "RECURSOS")
        self.assertEqual(filas[inicio + 1],
                         ("Recurso", "Cantidad", "Unidad", "Precio unitario ($)", "Total ($)"))
        return filas[inicio + 2:]

    def test_recursos_por_tarea_y_resumen_con_importes_en_pesos(self):
        resultado = self._resultado()
        wb = self._workbook([resultado, resultado])
        esperado = [
            ("Fabricación", 6, "Operarios-día", 1000, 6000),
            ("Montaje", 6, "Operarios-día", 2000, 12000),
            ("Grua", 4, "Días", 5000, 20000),
            ("Director de obra", 5, "Días", 3000, 15000),
            ("Técnico H y S", 2, "Días", 2500, 5000),
            ("TOTAL MANO DE OBRA — FABRICACIÓN", 6, "Operarios-día", None, 6000),
            ("TOTAL RECURSOS", None, None, None, 58000),
        ]
        for nombre in ("Tarea 1", "Tarea 2"):
            self.assertEqual(self._recursos(wb[nombre]), esperado)
        self.assertEqual(self._recursos(wb["Resumen"]), [
            (nombre, cantidad * 2, unidad, tarifa, total * 2)
            if cantidad is not None else (nombre, None, None, None, total * 2)
            for nombre, cantidad, unidad, tarifa, total in esperado
        ])
        for ws in wb:
            inicio = next(c.row for c in ws["A"] if c.value == "RECURSOS")
            fila = next(c.row for c in ws["A"]
                        if c.row > inicio and c.value == "Director de obra")
            self.assertEqual(ws.cell(fila, 4).data_type, "n")
            self.assertIn("$", ws.cell(fila, 4).number_format)
            self.assertIn("$", ws.cell(fila, 5).number_format)

    def test_tarifas_distintas_se_muestran_separadas(self):
        wb = self._workbook([self._resultado(), self._resultado(1200)])
        filas = self._recursos(wb["Resumen"])
        self.assertEqual([f for f in filas if f[0] == "Fabricación"], [
            ("Fabricación", 6, "Operarios-día", 1000, 6000),
            ("Fabricación", 6, "Operarios-día", 1200, 7200),
        ])
        self.assertEqual(filas[-1][-1], 117200)
        self.assertEqual(filas[-2],
                         ("TOTAL MANO DE OBRA — FABRICACIÓN", 12, "Operarios-día", None, 13200))

    def test_sin_tareas_y_dias_cero(self):
        wb = self._workbook([])
        self.assertEqual(self._recursos(wb["Resumen"])[0][0], "Sin recursos cargados.")
        pcts = {"gg_pct": 0, "beneficio_pct": 0, "imp_pct": 0}
        resultado = calcular_tarea([], pcts, [
            {"rubro": "ingeniero", "datos": {"dias": 0, "tarifa_dia": 3000}},
        ], pcts)
        wb = self._workbook([resultado])
        self.assertEqual(self._recursos(wb["Resumen"])[0],
                         ("Director de obra", 0, "Días", 3000, 0))

    def test_total_fabricacion_104_operarios_dia_desde_base_de_datos(self):
        ensure_tablas_presupuestos(self.db)
        presupuesto_id = crear_presupuesto(self.db, "Prueba", tipo_cambio_referencia=1500)
        for nombre, operarios, dias, tarifa in [
            ("Estructura", 3, 20, 127000),
            ("Cubierta", 4, 11, 130000),
        ]:
            tarea_id = crear_tarea(self.db, presupuesto_id, nombre)
            secciones = crear_secciones_tarea(self.db, tarea_id, obtener_config(self.db))
            for rubro in ("mano_obra", "consumibles"):
                crear_item_costo(self.db, secciones["FABRICACION"], rubro, {
                    "operarios": operarios, "dias": dias,
                    "tarifa_dh": tarifa if rubro == "mano_obra" else 500,
                })
            crear_item_costo(self.db, secciones["MONTAJE"], "mano_obra", {
                "operarios": 2, "dias": 3, "tarifa_dh": 190000,
            })
            crear_item_costo(self.db, secciones["MONTAJE"], "ingeniero", {
                "dias": 0.5, "tarifa_dia": 160000,
            })
            crear_item_costo(self.db, secciones["MONTAJE"], "tecnico_hys", {
                "dias": 0.5, "tarifa_dia": 160000,
            })

        wb = load_workbook(generar_reporte_explosion_insumos(self.db, presupuesto_id))
        self.addCleanup(wb.close)
        filas = self._recursos(wb["Resumen"])
        self.assertEqual(next(f for f in filas if f[0] == "TOTAL MANO DE OBRA — FABRICACIÓN"),
                         ("TOTAL MANO DE OBRA — FABRICACIÓN", 104, "Operarios-día",
                          None, 13340000))
        self.assertEqual(next(f for f in filas if f[0] == "Montaje"),
                         ("Montaje", 12, "Operarios-día", 190000, 2280000))
        for nombre in ("Director de obra", "Técnico H y S"):
            self.assertEqual(next(f for f in filas if f[0] == nombre),
                             (nombre, 1, "Días", 160000, 160000))
        self.assertEqual(filas[-1][-1], 15940000)
        for nombre, esperado in (("Estructura", 60), ("Cubierta", 44)):
            filas_tarea = self._recursos(wb[nombre])
            self.assertEqual(next(f[1] for f in filas_tarea
                                  if f[0] == "TOTAL MANO DE OBRA — FABRICACIÓN"), esperado)
            self.assertEqual(next(f for f in filas_tarea if f[0] == "Director de obra"),
                             ("Director de obra", 0.5, "Días", 160000, 80000))


if __name__ == "__main__":
    unittest.main()
